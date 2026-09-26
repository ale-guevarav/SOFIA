"""
GENERADOR DE REPORTES DE MÉTRICAS - SOFIA

Lee el JSON que genera metrics_logger.py y crea:

    metrics/reportes/<sesion>/
        reporte.xlsx      -> Excel con tablas y gráficas nativas
        reporte.md        -> Markdown con tablas e imágenes
        graficas/*.png    -> Gráficas hechas con matplotlib

Uso:
    python generate_report.py                       (sesión más reciente)
    python generate_report.py metrics/sesion_X.json (sesión específica)
"""

import glob
import json
import os
import sys

import matplotlib
matplotlib.use("Agg")  # Generar imágenes sin abrir ventanas
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# -----------------------------
# ETIQUETAS Y COLORES
# -----------------------------

OUTCOME_LABELS = {
    "acierto": "Acierto",
    "no_detectado": "No detectado",
    "confundido": "Confundido",
    "falso_positivo": "Falso positivo",
}

OUTCOME_COLORS = {
    "acierto": "#2ECC71",
    "no_detectado": "#95A5A6",
    "confundido": "#E67E22",
    "falso_positivo": "#E74C3C",
}

STAGE_LABELS = {
    "pose": "Pose Landmarker",
    "manos": "Hand Landmarker",
    "gestos": "Lógica de gestos",
    "dibujo": "Dibujo / interfaz",
}

STAGE_COLORS = ["#3498DB", "#9B59B6", "#1ABC9C", "#F1C40F", "#E67E22"]

BLUE = "#1F4E78"


# =============================================
# UTILIDADES
# =============================================

def fmt(value, suffix="", decimals=2):
    """Formato para texto: None -> guion."""
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.{decimals}f}{suffix}"
    return f"{value}{suffix}"


def stage_name(key):
    return STAGE_LABELS.get(key, key)


def latest_session(folder="metrics"):
    files = glob.glob(os.path.join(folder, "sesion_*.json"))
    if not files:
        return None
    return max(files, key=os.path.getmtime)


def gestures_with_trials(data):
    return [
        label for label, info in data["por_gesto"].items()
        if info["intentos"] > 0
    ]


# =============================================
# GRÁFICAS PNG (matplotlib)
# =============================================

def _save_figure(fig, folder, name):
    fig.tight_layout()
    fig.savefig(os.path.join(folder, name), dpi=150)
    plt.close(fig)
    return f"graficas/{name}"


def make_charts(data, charts_dir):

    os.makedirs(charts_dir, exist_ok=True)
    plt.rcParams.update({"font.size": 11, "axes.spines.top": False,
                         "axes.spines.right": False})
    charts = {}

    # -----------------------------
    # 1. DONA DE RESULTADOS
    # -----------------------------

    outcomes = {k: v for k, v in data["resultados"].items() if v > 0}

    if outcomes:
        fig, ax = plt.subplots(figsize=(6, 6))
        ax.pie(
            outcomes.values(),
            labels=[OUTCOME_LABELS[k] for k in outcomes],
            colors=[OUTCOME_COLORS[k] for k in outcomes],
            autopct=lambda p: f"{p:.1f}%" if p > 0 else "",
            startangle=90,
            pctdistance=0.78,
            wedgeprops={"width": 0.45, "edgecolor": "white", "linewidth": 2},
        )
        rate = data["resumen"]["tasa_acierto_global"]
        ax.text(0, 0, f"{fmt(rate, '%', 1)}\nacierto", ha="center",
                va="center", fontsize=18, fontweight="bold")
        ax.set_title("Resultados de los intentos", fontweight="bold")
        charts["resultados"] = _save_figure(fig, charts_dir, "resultados.png")

    # -----------------------------
    # 2. TASA DE ACIERTO POR GESTO
    # -----------------------------

    labels = gestures_with_trials(data)

    if labels:
        rates = [data["por_gesto"][g]["tasa_acierto"] or 0 for g in labels]
        colors = ["#2ECC71" if r >= 80 else "#F1C40F" if r >= 60 else "#E74C3C"
                  for r in rates]

        fig, ax = plt.subplots(figsize=(8, 4.5))
        bars = ax.bar(labels, rates, color=colors)
        ax.bar_label(bars, labels=[f"{r:.0f}%" for r in rates], padding=3)
        ax.axhline(80, color="gray", linestyle="--", linewidth=1)
        ax.text(len(labels) - 0.5, 81, "meta 80%", color="gray",
                ha="right", fontsize=9)
        ax.set_ylim(0, 110)
        ax.set_ylabel("Tasa de acierto (%)")
        ax.set_title("Tasa de acierto por gesto", fontweight="bold")
        charts["acierto_por_gesto"] = _save_figure(fig, charts_dir, "acierto_por_gesto.png")

    # -----------------------------
    # 3. FPS Y LATENCIA EN EL TIEMPO
    # -----------------------------

    timeline = data["linea_de_tiempo"]

    if timeline:
        seconds = [p["segundo"] for p in timeline]
        fps = [p["fps"] for p in timeline]
        latency = [p["latencia_ms"] for p in timeline]

        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 6), sharex=True)

        ax1.plot(seconds, fps, color="#3498DB", linewidth=1.8)
        ax1.axhline(data["fps"]["promedio_real"], color="#E74C3C",
                    linestyle="--", label=f"Promedio: {data['fps']['promedio_real']:.1f}")
        ax1.fill_between(seconds, fps, alpha=0.15, color="#3498DB")
        ax1.set_ylabel("FPS")
        ax1.set_title("FPS y latencia durante la sesión", fontweight="bold")
        ax1.legend(loc="lower right")

        ax2.plot(seconds, latency, color="#9B59B6", linewidth=1.8)
        ax2.axhline(data["latencia_procesamiento_ms"]["promedio"], color="#E74C3C",
                    linestyle="--",
                    label=f"Promedio: {data['latencia_procesamiento_ms']['promedio']:.1f} ms")
        ax2.fill_between(seconds, latency, alpha=0.15, color="#9B59B6")
        ax2.set_ylabel("Latencia (ms)")
        ax2.set_xlabel("Tiempo de sesión (s)")
        ax2.legend(loc="upper right")

        charts["linea_de_tiempo"] = _save_figure(fig, charts_dir, "linea_de_tiempo.png")

    # -----------------------------
    # 4. PASTEL DE ETAPAS
    # -----------------------------

    stages = data["etapas_ms"]

    if stages:
        names = list(stages.keys())
        fig, ax = plt.subplots(figsize=(6.5, 6))
        ax.pie(
            [stages[n]["promedio"] for n in names],
            labels=[f"{stage_name(n)}\n{stages[n]['promedio']:.1f} ms" for n in names],
            colors=STAGE_COLORS[:len(names)],
            autopct="%1.1f%%",
            startangle=90,
            wedgeprops={"edgecolor": "white", "linewidth": 2},
        )
        ax.set_title("¿En qué se va el tiempo de cada frame?", fontweight="bold")
        charts["etapas"] = _save_figure(fig, charts_dir, "etapas.png")

    # -----------------------------
    # 5. MATRIZ DE CONFUSIÓN
    # -----------------------------

    if data["intentos"]:
        labels_all = data["matriz_confusion"]["etiquetas"]
        values = data["matriz_confusion"]["valores"]
        matrix = [[values[e][p] for p in labels_all] for e in labels_all]
        max_value = max(max(row) for row in matrix) or 1

        fig, ax = plt.subplots(figsize=(7, 6))
        image = ax.imshow(matrix, cmap="Blues")
        ax.set_xticks(range(len(labels_all)), labels_all, rotation=45, ha="right")
        ax.set_yticks(range(len(labels_all)), labels_all)
        ax.set_xlabel("Detectado por el sistema")
        ax.set_ylabel("Gesto esperado")
        ax.set_title("Matriz de confusión", fontweight="bold")

        for i, row in enumerate(matrix):
            for j, value in enumerate(row):
                ax.text(j, i, value, ha="center", va="center",
                        color="white" if value > max_value / 2 else "black")

        fig.colorbar(image, ax=ax, fraction=0.046)
        charts["matriz"] = _save_figure(fig, charts_dir, "matriz_confusion.png")

    # -----------------------------
    # 6. TIEMPO DE RESPUESTA POR GESTO
    # -----------------------------

    rt_labels = [
        g for g in labels
        if data["por_gesto"][g]["tiempo_respuesta_s"]["n"] > 0
    ]

    if rt_labels:
        means = [data["por_gesto"][g]["tiempo_respuesta_s"]["promedio"] for g in rt_labels]
        stds = [data["por_gesto"][g]["tiempo_respuesta_s"]["desv_std"] for g in rt_labels]

        fig, ax = plt.subplots(figsize=(8, 4.5))
        bars = ax.bar(rt_labels, means, yerr=stds, capsize=6, color="#1ABC9C")
        ax.bar_label(bars, labels=[f"{m:.2f}s" for m in means], padding=3)
        ax.set_ylabel("Segundos")
        ax.set_title("Tiempo de reconocimiento promedio (± desv. std)",
                     fontweight="bold")
        charts["tiempo_respuesta"] = _save_figure(fig, charts_dir, "tiempo_respuesta.png")

    return charts


# =============================================
# EXCEL (openpyxl)
# =============================================

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
SUBTITLE_FONT = Font(bold=True, size=12, color="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ZEBRA_FILL = PatternFill("solid", fgColor="EAF1FB")


def write_table(ws, row, col, headers, rows, number_format="0.00"):
    """Escribe una tabla con estilo. Regresa (fila_encabezado, última_fila)."""

    for j, header in enumerate(headers):
        cell = ws.cell(row=row, column=col + j, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, values in enumerate(rows, start=1):
        for j, value in enumerate(values):
            cell = ws.cell(row=row + i, column=col + j, value=value)
            cell.border = BORDER
            if isinstance(value, float):
                cell.number_format = number_format
            if i % 2 == 0:
                cell.fill = ZEBRA_FILL

    return row, row + len(rows)


def autosize(ws, min_width=10, max_width=45):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                length = len(str(cell.value))
                widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = (
            max(min_width, min(max_width, width + 2))
        )


def _show_axes(chart):
    # En Excel reciente los ejes pueden salir ocultos si no se indica
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def build_excel(data, path):

    wb = Workbook()
    summary = data["resumen"]

    # -----------------------------
    # HOJA 1: RESUMEN
    # -----------------------------

    ws = wb.active
    ws.title = "Resumen"

    ws["A1"] = "SOFIA - Métricas de desempeño"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Sesión: {data['sesion']['nombre']}   |   {data['sesion']['inicio']} → {data['sesion']['fin']}"

    kpis = [
        ("Tasa de acierto global (%)", summary["tasa_acierto_global"]),
        ("Reconocimiento de gestos (%)", summary["tasa_reconocimiento_gestos"]),
        ("Rechazo correcto - NINGUNO (%)", summary["tasa_rechazo_correcto"]),
        ("Tiempo de respuesta promedio (s)", summary["tiempo_respuesta_promedio_s"]),
        ("FPS promedio (real)", summary["fps_promedio"]),
        ("Capacidad de procesamiento (FPS)", data["fps"]["capacidad_procesamiento"]),
        ("Latencia promedio por frame (ms)", summary["latencia_promedio_ms"]),
        ("Latencia p95 por frame (ms)", data["latencia_procesamiento_ms"]["p95"]),
        ("Total de intentos", summary["total_intentos"]),
        ("Aciertos", summary["aciertos"]),
        ("Frames procesados", summary["frames_procesados"]),
        ("Frames con persona (%)", summary["porcentaje_frames_con_persona"]),
        ("Duración de la sesión (s)", summary["duracion_s"]),
    ]

    write_table(ws, 4, 1, ["Métrica", "Valor"], kpis)

    # Resaltar las 3 métricas principales
    for r in (5, 8, 9):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).font = Font(bold=True, color="1F4E78")

    # Datos para el pastel de resultados
    outcome_row = 4 + len(kpis) + 3
    ws.cell(row=outcome_row - 1, column=1, value="Resultados de los intentos").font = SUBTITLE_FONT
    outcome_rows = [(OUTCOME_LABELS[k], v) for k, v in data["resultados"].items()]
    header, last = write_table(ws, outcome_row, 1, ["Resultado", "Cantidad"], outcome_rows)

    if summary["total_intentos"]:
        pie = PieChart()
        pie.title = "Resultados de los intentos"
        pie.add_data(Reference(ws, min_col=2, min_row=header, max_row=last), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=header + 1, max_row=last))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.height, pie.width = 8, 12
        ws.add_chart(pie, "D4")

    # Datos para el pastel de etapas
    stage_row = last + 3
    ws.cell(row=stage_row - 1, column=1, value="Tiempo promedio por etapa").font = SUBTITLE_FONT
    stage_rows = [
        (stage_name(k), v["promedio"]) for k, v in data["etapas_ms"].items()
    ]
    header, last = write_table(ws, stage_row, 1, ["Etapa", "ms promedio"], stage_rows)

    if stage_rows:
        pie2 = PieChart()
        pie2.title = "Distribución del tiempo por frame"
        pie2.add_data(Reference(ws, min_col=2, min_row=header, max_row=last), titles_from_data=True)
        pie2.set_categories(Reference(ws, min_col=1, min_row=header + 1, max_row=last))
        pie2.dataLabels = DataLabelList()
        pie2.dataLabels.showPercent = True
        pie2.height, pie2.width = 8, 12
        ws.add_chart(pie2, "D21")

    autosize(ws)

    # -----------------------------
    # HOJA 2: POR GESTO
    # -----------------------------

    ws = wb.create_sheet("Por gesto")
    ws["A1"] = "Desempeño por gesto"
    ws["A1"].font = TITLE_FONT

    headers = ["Gesto", "Intentos", "Aciertos", "No detectado", "Confundido",
               "Falso positivo", "Tasa acierto (%)", "Precisión (%)",
               "Recall (%)", "F1 (%)", "T. respuesta prom (s)",
               "T. respuesta p95 (s)"]

    rows = []
    for label, info in data["por_gesto"].items():
        rt = info["tiempo_respuesta_s"]
        rows.append([
            label, info["intentos"], info["aciertos"], info["no_detectado"],
            info["confundido"], info["falso_positivo"], info["tasa_acierto"],
            info["precision"], info["recall"], info["f1"],
            rt["promedio"] if rt["n"] else None,
            rt["p95"] if rt["n"] else None,
        ])

    header, last = write_table(ws, 3, 1, headers, rows)

    bar = BarChart()
    bar.type = "col"
    bar.title = "Tasa de acierto por gesto (%)"
    bar.y_axis.title = "%"
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 100
    bar.add_data(Reference(ws, min_col=7, min_row=header, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=header + 1, max_row=last))
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.legend = None
    bar.height, bar.width = 8, 16
    _show_axes(bar)
    ws.add_chart(bar, "A12")

    bar2 = BarChart()
    bar2.type = "col"
    bar2.title = "Tiempo de respuesta promedio (s)"
    bar2.add_data(Reference(ws, min_col=11, min_row=header, max_row=last), titles_from_data=True)
    bar2.set_categories(Reference(ws, min_col=1, min_row=header + 1, max_row=last))
    bar2.dataLabels = DataLabelList()
    bar2.dataLabels.showVal = True
    bar2.legend = None
    bar2.height, bar2.width = 8, 16
    _show_axes(bar2)
    ws.add_chart(bar2, "H12")

    autosize(ws)

    # -----------------------------
    # HOJA 3: RENDIMIENTO
    # -----------------------------

    ws = wb.create_sheet("Rendimiento")
    ws["A1"] = "Rendimiento: FPS y latencia"
    ws["A1"].font = TITLE_FONT

    stats_keys = ["promedio", "mediana", "minimo", "maximo", "p95", "desv_std"]
    fps_inst = data["fps"]["instantaneo"]
    latency = data["latencia_procesamiento_ms"]

    stats_rows = [
        [key.replace("_", " ").capitalize(), fps_inst[key], latency[key]]
        for key in stats_keys
    ]
    write_table(ws, 3, 1, ["Estadístico", "FPS instantáneo", "Latencia (ms)"], stats_rows)

    row = 3 + len(stats_rows) + 3
    ws.cell(row=row - 1, column=1, value="Latencia por etapa (ms)").font = SUBTITLE_FONT
    stage_rows = [
        [stage_name(k), v["promedio"], v["p95"], v["maximo"], v["porcentaje"]]
        for k, v in data["etapas_ms"].items()
    ]
    write_table(ws, row, 1, ["Etapa", "Promedio", "p95", "Máximo", "% del frame"], stage_rows)

    # Línea de tiempo (columnas H-J)
    timeline_rows = [
        [p["segundo"], p["fps"], p["latencia_ms"]] for p in data["linea_de_tiempo"]
    ]
    ws.cell(row=2, column=8, value="Línea de tiempo (1 punto por segundo)").font = SUBTITLE_FONT
    header, last = write_table(ws, 3, 8, ["Segundo", "FPS", "Latencia (ms)"], timeline_rows)

    if timeline_rows:
        categories = Reference(ws, min_col=8, min_row=header + 1, max_row=last)

        line = LineChart()
        line.title = "FPS durante la sesión"
        line.y_axis.title = "FPS"
        line.x_axis.title = "Segundo"
        line.add_data(Reference(ws, min_col=9, min_row=header, max_row=last), titles_from_data=True)
        line.set_categories(categories)
        line.height, line.width = 8, 18
        _show_axes(line)
        ws.add_chart(line, "L3")

        line2 = LineChart()
        line2.title = "Latencia de procesamiento"
        line2.y_axis.title = "ms"
        line2.x_axis.title = "Segundo"
        line2.add_data(Reference(ws, min_col=10, min_row=header, max_row=last), titles_from_data=True)
        line2.set_categories(categories)
        line2.height, line2.width = 8, 18
        _show_axes(line2)
        ws.add_chart(line2, "L20")

    autosize(ws)

    # -----------------------------
    # HOJA 4: MATRIZ DE CONFUSIÓN
    # -----------------------------

    ws = wb.create_sheet("Matriz de confusion")
    ws["A1"] = "Matriz de confusión (filas = esperado, columnas = detectado)"
    ws["A1"].font = TITLE_FONT

    labels_all = data["matriz_confusion"]["etiquetas"]
    values = data["matriz_confusion"]["valores"]
    matrix_rows = [[e] + [values[e][p] for p in labels_all] for e in labels_all]
    header, last = write_table(ws, 3, 1, ["Esperado \\ Detectado"] + labels_all, matrix_rows)

    first_cell = f"B{header + 1}"
    last_cell = f"{get_column_letter(1 + len(labels_all))}{last}"
    ws.conditional_formatting.add(
        f"{first_cell}:{last_cell}",
        ColorScaleRule(start_type="min", start_color="FFFFFF",
                       end_type="max", end_color="2E75B6"),
    )

    for r in range(header + 1, last + 1):
        for c in range(2, 2 + len(labels_all)):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")

    ws.cell(row=last + 2, column=1,
            value="La diagonal son los aciertos. Fuera de la diagonal están las confusiones.")
    autosize(ws)

    # -----------------------------
    # HOJA 5: INTENTOS
    # -----------------------------

    ws = wb.create_sheet("Intentos")
    ws["A1"] = "Registro de intentos"
    ws["A1"].font = TITLE_FONT

    trial_rows = [
        [t["id"], t["momento_s"], t["esperado"], t["detectado"],
         OUTCOME_LABELS[t["resultado"]], "Sí" if t["correcto"] else "No",
         t["tiempo_respuesta_s"]]
        for t in data["intentos"]
    ]
    write_table(ws, 3, 1, ["#", "Momento (s)", "Esperado", "Detectado",
                           "Resultado", "Correcto", "T. respuesta (s)"], trial_rows)

    # Colorear columna de resultado
    for r in range(4, 4 + len(trial_rows)):
        cell = ws.cell(row=r, column=6)
        cell.font = Font(bold=True, color="1E8449" if cell.value == "Sí" else "C0392B")

    ws.freeze_panes = "A4"
    autosize(ws)

    wb.save(path)


# =============================================
# MARKDOWN
# =============================================

def md_table(headers, rows):
    lines = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join([" --- "] * len(headers)) + "|",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(v) for v in row) + " |")
    return "\n".join(lines)


def build_markdown(data, charts, path):

    s = data["resumen"]
    session = data["sesion"]
    out = []

    out.append("# SOFIA — Reporte de métricas de desempeño\n")
    out.append(f"**Sesión:** `{session['nombre']}`  ")
    out.append(f"**Inicio:** {session['inicio']} — **Fin:** {session['fin']}  ")
    out.append(f"**Equipo:** {session['sistema'].get('so', '—')} · "
               f"Python {session['sistema'].get('python', '—')} · "
               f"OpenCV {session['sistema'].get('opencv', '—')}\n")

    # Métricas principales
    out.append("## Métricas principales\n")
    out.append(md_table(["Métrica", "Valor"], [
        ["🎯 **Tasa de acierto global**", f"**{fmt(s['tasa_acierto_global'], '%')}**"],
        ["Reconocimiento de gestos", fmt(s["tasa_reconocimiento_gestos"], "%")],
        ["Rechazo correcto (NINGUNO)", fmt(s["tasa_rechazo_correcto"], "%")],
        ["⏱️ **Tiempo de respuesta promedio**", f"**{fmt(s['tiempo_respuesta_promedio_s'], ' s')}**"],
        ["🎞️ **FPS promedio**", f"**{fmt(s['fps_promedio'])}**"],
        ["Capacidad de procesamiento", fmt(data["fps"]["capacidad_procesamiento"], " FPS")],
        ["Latencia promedio por frame", fmt(s["latencia_promedio_ms"], " ms")],
        ["Latencia p95 por frame", fmt(data["latencia_procesamiento_ms"]["p95"], " ms")],
        ["Intentos / aciertos", f"{s['total_intentos']} / {s['aciertos']}"],
        ["Frames procesados", s["frames_procesados"]],
        ["Frames con persona", fmt(s["porcentaje_frames_con_persona"], "%")],
        ["Duración", fmt(s["duracion_s"], " s")],
    ]))
    out.append("")

    if "resultados" in charts:
        out.append(f"![Resultados]({charts['resultados']})\n")

    # Por gesto
    out.append("## Desempeño por gesto\n")
    rows = []
    for label, info in data["por_gesto"].items():
        rt = info["tiempo_respuesta_s"]
        rows.append([
            f"**{label}**", info["intentos"], info["aciertos"],
            fmt(info["tasa_acierto"], "%"), fmt(info["precision"], "%"),
            fmt(info["recall"], "%"), fmt(info["f1"], "%"),
            fmt(rt["promedio"], " s") if rt["n"] else "—",
        ])
    out.append(md_table(["Gesto", "Intentos", "Aciertos", "Acierto", "Precisión",
                         "Recall", "F1", "T. respuesta"], rows))
    out.append("")

    for key, title in (("acierto_por_gesto", "Acierto por gesto"),
                       ("tiempo_respuesta", "Tiempo de respuesta"),
                       ("matriz", "Matriz de confusión")):
        if key in charts:
            out.append(f"![{title}]({charts[key]})\n")

    # Rendimiento
    out.append("## Rendimiento (FPS y latencia)\n")
    fps_inst = data["fps"]["instantaneo"]
    lat = data["latencia_procesamiento_ms"]
    out.append(md_table(["Estadístico", "FPS instantáneo", "Latencia (ms)"], [
        [k.replace("_", " ").capitalize(), fmt(fps_inst[k]), fmt(lat[k])]
        for k in ["promedio", "mediana", "minimo", "maximo", "p95", "desv_std"]
    ]))
    out.append("")

    if "linea_de_tiempo" in charts:
        out.append(f"![Línea de tiempo]({charts['linea_de_tiempo']})\n")

    out.append("### Latencia por etapa\n")
    out.append(md_table(["Etapa", "Promedio (ms)", "p95 (ms)", "% del frame"], [
        [stage_name(k), fmt(v["promedio"]), fmt(v["p95"]), fmt(v["porcentaje"], "%")]
        for k, v in data["etapas_ms"].items()
    ]))
    out.append("")

    if "etapas" in charts:
        out.append(f"![Etapas]({charts['etapas']})\n")

    # Intentos
    out.append("## Registro de intentos\n")
    if data["intentos"]:
        out.append(md_table(["#", "Esperado", "Detectado", "Resultado", "T. respuesta"], [
            [t["id"], t["esperado"], t["detectado"],
             ("✅ " if t["correcto"] else "❌ ") + OUTCOME_LABELS[t["resultado"]],
             fmt(t["tiempo_respuesta_s"], " s")]
            for t in data["intentos"]
        ]))
    else:
        out.append("_No se realizaron intentos en esta sesión._")
    out.append("")

    # Glosario
    out.append("## ¿Cómo se calculan?\n")
    out.append("- **Tasa de acierto:** intentos correctos / intentos totales × 100.")
    out.append("- **Precisión:** de las veces que el sistema dijo *X*, cuántas eran realmente *X*.")
    out.append("- **Recall:** de las veces que la persona hizo *X*, cuántas detectó el sistema.")
    out.append("- **Tiempo de respuesta:** segundos desde la señal *YA* hasta que el sistema reporta el gesto (solo aciertos).")
    out.append("- **FPS promedio:** frames totales / tiempo total del ciclo.")
    out.append("- **Latencia por frame:** tiempo de procesamiento (pose + manos + gestos + dibujo), sin contar la espera de la cámara.")

    with open(path, "w", encoding="utf-8") as file:
        file.write("\n".join(out))


# =============================================
# FUNCIÓN PRINCIPAL
# =============================================

def generate(json_path):

    with open(json_path, encoding="utf-8") as file:
        data = json.load(file)

    session_name = data["sesion"]["nombre"]
    out_dir = os.path.join(os.path.dirname(json_path) or ".", "reportes", session_name)
    os.makedirs(out_dir, exist_ok=True)

    charts = make_charts(data, os.path.join(out_dir, "graficas"))

    xlsx_path = os.path.join(out_dir, "reporte.xlsx")
    md_path = os.path.join(out_dir, "reporte.md")

    build_excel(data, xlsx_path)
    build_markdown(data, charts, md_path)

    print(f"[REPORTE] Excel:    {xlsx_path}")
    print(f"[REPORTE] Markdown: {md_path}")

    return xlsx_path, md_path


if __name__ == "__main__":

    path = sys.argv[1] if len(sys.argv) > 1 else latest_session()

    if path is None:
        print("No se encontró ningún archivo metrics/sesion_*.json")
        sys.exit(1)

    generate(path)